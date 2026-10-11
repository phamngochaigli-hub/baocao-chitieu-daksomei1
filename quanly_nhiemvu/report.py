"""Xuất báo cáo ra tệp Word (.docx) theo thể thức văn bản hành chính."""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .models import BaoCao, NhiemVu
from .workflow import muc_canh_bao


def _doan(doc, chu: str, dam: bool = False, can_giua: bool = False, co: int = 13):
    p = doc.add_paragraph()
    run = p.add_run(chu)
    run.bold = dam
    run.font.size = Pt(co)
    if can_giua:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def xuat_docx(bc: BaoCao, ds: list[NhiemVu], ten_co_quan: str, hom_nay: date) -> bytes:
    doc = Document()
    kieu = doc.styles["Normal"]
    kieu.font.name = "Times New Roman"
    kieu.font.size = Pt(13)
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)

    dau = doc.add_table(rows=1, cols=2)
    trai, phai = dau.rows[0].cells
    trai.paragraphs[0].add_run(ten_co_quan.upper()).bold = True
    trai.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = phai.paragraphs[0]
    p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n").bold = True
    p.add_run("Độc lập - Tự do - Hạnh phúc").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = phai.add_paragraph(f"ngày {hom_nay.day:02d} tháng {hom_nay.month:02d} năm {hom_nay.year}")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.runs[0].italic = True

    doc.add_paragraph()
    _doan(doc, "BÁO CÁO", dam=True, can_giua=True, co=14)
    _doan(doc, bc.tieu_de, dam=True, can_giua=True)
    _doan(doc, f"(Từ ngày {bc.tu_ngay:%d/%m/%Y} đến ngày {bc.den_ngay:%d/%m/%Y})", can_giua=True)
    if bc.don_vi:
        _doan(doc, f"Đơn vị: {bc.don_vi}", can_giua=True)

    for dong in bc.noi_dung.splitlines():
        dong = dong.rstrip()
        la_tieu_de = dong[:4].strip().rstrip(".") in {"I", "II", "III", "IV", "V", "VI"}
        p = _doan(doc, dong, dam=la_tieu_de)
        p.paragraph_format.first_line_indent = None if la_tieu_de else Cm(1)

    doc.add_page_break()
    _doan(doc, "PHỤ LỤC: DANH MỤC NHIỆM VỤ", dam=True, can_giua=True)
    bang = doc.add_table(rows=1, cols=6)
    bang.style = "Table Grid"
    bang.alignment = WD_TABLE_ALIGNMENT.CENTER
    for o, ten in zip(bang.rows[0].cells, ["TT", "Nội dung nhiệm vụ", "Đơn vị chủ trì", "Hạn", "Tiến độ", "Tình trạng"]):
        o.text = ""
        o.paragraphs[0].add_run(ten).bold = True
    for i, nv in enumerate(ds, 1):
        o = bang.add_row().cells
        o[0].text = str(i)
        o[1].text = nv.noi_dung
        o[2].text = nv.don_vi_chu_tri
        o[3].text = nv.han_xu_ly.strftime("%d/%m/%Y") if nv.han_xu_ly else ""
        o[4].text = f"{nv.tien_do}%"
        o[5].text = muc_canh_bao(nv, hom_nay)
    for hang in bang.rows:
        for o in hang.cells:
            for p in o.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)

    doc.add_paragraph()
    ky = doc.add_paragraph()
    ky.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ky.add_run("NGƯỜI LẬP\n").bold = True
    ky.add_run(bc.nguoi_lap)
    if bc.nguoi_duyet:
        ky.add_run("\n\nNGƯỜI DUYỆT\n").bold = True
        ky.add_run(bc.nguoi_duyet)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
